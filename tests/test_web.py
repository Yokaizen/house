from datetime import date, timedelta

from openpyxl import load_workbook
from io import BytesIO

from app.models import Audit, Pointage, Utilisateur
from tests.conftest import SITE


def test_connexion_obligatoire(client):
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/connexion"


def test_connexion_mauvais_mdp(client, donnees):
    r = client.post("/connexion", data={"identifiant": "SAL01", "mot_de_passe": "faux"})
    assert r.status_code == 400 and "incorrect" in r.text


def test_badger_ne_stocke_pas_la_position(client_de, db, donnees):
    c = client_de("sal")
    r = c.post("/api/badger", json={"type": "entree", "latitude": SITE[0] + 0.0003, "longitude": SITE[1],
                                    "precision": 12, "appareil": "mobile"},
               headers={"X-CSRF-Token": c.csrf})
    assert r.status_code == 200, r.text
    assert r.json()["zone"] == "dans_zone"
    assert "latitude" not in r.json()
    p = db.query(Pointage).one()
    assert p.zone_statut == "dans_zone" and p.appareil == "mobile" and p.precision_m == 12
    colonnes = {c.key for c in Pointage.__table__.columns}
    assert not colonnes & {"latitude", "longitude", "lat", "lon"}

    # Le départ depuis le domicile est aussi dans la zone.
    r = c.post("/api/badger", json={"type": "sortie", "latitude": 48.80, "longitude": 2.20, "precision": 30},
               headers={"X-CSRF-Token": c.csrf, "User-Agent": "Mozilla/5.0 (Windows NT 10.0)"})
    assert r.json()["zone"] == "dans_zone"
    # Journée terminée (une seule paire par défaut).
    r = c.post("/api/badger", json={"type": "entree"}, headers={"X-CSRF-Token": c.csrf})
    assert r.status_code == 409


def test_badger_sans_position_et_sans_csrf(client_de, db, donnees):
    c = client_de("autre")
    assert c.post("/api/badger", json={"type": "entree"}).status_code == 403
    r = c.post("/api/badger", json={"type": "entree"}, headers={"X-CSRF-Token": c.csrf})
    assert r.json()["zone"] == "indisponible"


def test_oubli(client_de, db, donnees):
    c = client_de("sal")
    hier = (date.today() - timedelta(days=1)).isoformat()
    r = c.post("/oubli", data={"csrf_token": c.csrf, "type": "entree", "jour": hier, "heure": "08:15",
                               "motif": "Batterie vide"}, follow_redirects=False)
    assert r.status_code == 303
    p = db.query(Pointage).one()
    assert p.source == "manuel" and p.motif == "Batterie vide"
    assert db.query(Audit).filter_by(action="saisie_manuelle").count() == 1
    page = c.get(f"/mes-badgeages?mois={hier[:7]}").text
    assert "Saisie manuelle" in page


def test_droits_de_visibilite(client_de, donnees):
    sal, autre = donnees["sal"], donnees["autre"]
    chef = client_de("chef")
    assert chef.get("/equipe").status_code == 200
    assert chef.get(f"/equipe/{sal.id}").status_code == 200
    assert chef.get(f"/equipe/{autre.id}").status_code == 403
    assert chef.get("/admin/utilisateurs").status_code == 403
    assert client_de("autre").get("/equipe").status_code == 403
    rh = client_de("rh")
    assert rh.get(f"/equipe/{autre.id}").status_code == 200
    assert "Corriger" not in rh.get(f"/equipe/{autre.id}?periode=mois").text


def test_correction_par_manager_tracee(client_de, db, donnees):
    sal = donnees["sal"]
    chef = client_de("chef")
    hier = (date.today() - timedelta(days=1)).isoformat()
    r = chef.post(f"/equipe/{sal.id}/ajout", data={"csrf_token": chef.csrf, "type": "entree", "jour": hier,
                                                   "heure": "07:00", "motif": "Badge oublié"}, follow_redirects=False)
    assert r.status_code == 303
    p = db.query(Pointage).one()
    assert p.source == "correction" and p.cree_par_id == donnees["chef"].id
    r = chef.post(f"/pointages/{p.id}/modifier", data={"csrf_token": chef.csrf, "type": "entree", "jour": hier,
                                                        "heure": "07:30", "motif": "Erreur d'heure"},
                  follow_redirects=False)
    assert r.status_code == 303
    r = chef.post(f"/pointages/{p.id}/supprimer", data={"csrf_token": chef.csrf, "motif": "Doublon"},
                  follow_redirects=False)
    assert r.status_code == 303
    actions = [a.action for a in db.query(Audit).order_by(Audit.id)]
    assert actions == ["ajout_pointage", "modification_pointage", "suppression_pointage"]
    # Un salarié ne peut pas corriger ses propres badgeages.
    soi = client_de("sal")
    assert soi.post(f"/equipe/{sal.id}/ajout", data={"csrf_token": soi.csrf, "type": "entree", "jour": hier,
                                                     "heure": "07:00", "motif": "x" * 5}).status_code == 403


def test_import_csv(client_de, db, donnees):
    admin = client_de("admin")
    csv = ("matricule;nom;prenom;email;etablissement;matricule_manager;role;heures_hebdo;teletravail;"
           "adresse_numero;adresse_voie;adresse_complement;code_postal;ville;date_sortie\n"
           "B0002;Petit;Léa;lea@proximy.net;STAINS;B0001;salarie;24,5;oui;3;rue Haute;;93240;Stains;\n"
           "B0001;Grand;Paul;paul@proximy.net;Stains;;manager;35;non;;;;;;\n"
           "B0003;Sans;;x@proximy.net;STAINS;;salarie;35;oui;;;;;;\n"
           "B0004;Inconnu;Ana;;NANTES;;salarie;35;oui;;;;;;\n").encode("cp1252")
    r = admin.post("/admin/import", data={"csrf_token": admin.csrf},
                   files={"fichier": ("u.csv", csv, "text/csv")})
    assert r.status_code == 200
    assert "Mots de passe provisoires" in r.text
    lea = db.query(Utilisateur).filter_by(matricule="B0002").one()
    paul = db.query(Utilisateur).filter_by(matricule="B0001").one()
    assert lea.manager_id == paul.id and float(lea.heures_hebdo) == 24.5 and lea.prenom == "Léa"
    assert paul.role == "manager" and not paul.teletravail_autorise
    assert db.query(Utilisateur).filter_by(matricule="B0003").count() == 0
    assert "NANTES" in r.text

    # Réimport : mise à jour sans doublon.
    r = admin.post("/admin/import", data={"csrf_token": admin.csrf},
                   files={"fichier": ("u.csv", b"matricule;nom;prenom;heures_hebdo\nB0002;Petit;Lea;28\n", "text/csv")})
    db.expire_all()
    assert float(db.query(Utilisateur).filter_by(matricule="B0002").one().heures_hebdo) == 28
    assert lea.manager_id == paul.id


def test_export_excel(client_de, donnees):
    admin = client_de("admin")
    r = admin.get("/equipe/export.xlsx?periode=3mois")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Synthèse", "Détail", "Légende"]


def test_export_excel_sans_formule(client_de):
    # Un motif saisi par un salarié ne doit pas devenir une formule dans le fichier ouvert par les RH.
    sal = client_de("sal")
    motif = '=HYPERLINK("https://exemple.invalid","Voir le justificatif")'
    hier = date.today() - timedelta(days=1)
    r = sal.post("/oubli", data={"type": "entree", "jour": hier.isoformat(), "heure": "08:00", "motif": motif,
                                 "csrf_token": sal.csrf})
    assert r.status_code == 200, r.text
    r = client_de("admin").get("/equipe/export.xlsx?periode=3mois")
    detail = load_workbook(BytesIO(r.content))["Détail"]
    cellules = [c for ligne in detail.iter_rows() for c in ligne if c.value == motif]
    assert cellules and all(c.data_type == "s" for c in cellules)


def test_pages_admin(client_de, donnees):
    admin = client_de("admin")
    for url in ["/admin/utilisateurs", "/admin/utilisateurs/nouveau", f"/admin/utilisateurs/{donnees['sal'].id}",
                "/admin/import", "/admin/etablissements", "/admin/etablissements/nouveau",
                f"/admin/etablissements/{donnees['etab'].id}", "/admin/parametres", "/admin/audit",
                "/equipe?periode=12mois", "/validations", "/confidentialite", "/oubli", "/",
                "/admin/import/modele.csv"]:
        assert admin.get(url).status_code == 200, url


def test_creation_utilisateur_et_premier_mot_de_passe(client_de, client, db, donnees):
    admin = client_de("admin")
    r = admin.post("/admin/utilisateurs/enregistrer", data={
        "csrf_token": admin.csrf, "matricule": "N001", "nom": "Neuf", "prenom": "Nina", "role": "salarie",
        "heures_hebdo": "35", "actif": "on", "teletravail_autorise": "on"})
    assert r.status_code == 200
    mdp = r.text.split('class="mdp-provisoire">')[1].split("<")[0]
    r = client.post("/connexion", data={"identifiant": "N001", "mot_de_passe": mdp}, follow_redirects=False)
    assert r.headers["location"] == "/mot-de-passe"
    assert client.get("/", follow_redirects=False).headers["location"] == "/mot-de-passe"


def test_parametres(client_de, db, donnees):
    admin = client_de("admin")
    from app.services import parametres as P
    data = {d.cle: str(d.defaut) for d in P.DEFINITIONS if d.type is not bool}
    data.update({"csrf_token": admin.csrf, "tous_sites_autorises": "on", "paires_max_par_jour": "2"})
    r = admin.post("/admin/parametres", data=data, follow_redirects=False)
    assert r.status_code == 303
    db.expire_all()
    assert P.lire(db, "paires_max_par_jour") == 2
    assert P.lire(db, "validation_manager") is False


def test_purge_compte_sorti(db, donnees):
    from app.jobs import purger
    autre = donnees["autre"]
    autre.date_sortie = date.today() - timedelta(days=120)
    db.add(Pointage(utilisateur_id=autre.id, type="entree", horodatage=__import__("datetime").datetime.now()))
    db.commit()
    purger(db)
    assert db.get(Utilisateur, autre.id) is None
    assert db.query(Pointage).count() == 0
