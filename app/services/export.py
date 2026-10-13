"""Export Excel (charte Proximy : en-têtes marine, titres Cambria, corps Calibri)."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import ZONE_LIBELLES
from app.services.pointages import ALERTES, Synthese
from app.services.temps import format_date, format_heure, libelle_mois

MARINE, VIOLET, ROSE, ENCRE = "061E66", "4B4584", "E80A65", "3A3F4B"
CARTE, FILET, GRIS = "F2F4F8", "E2E6EF", "7A8090"

_titre = Font(name="Cambria", size=16, bold=True, color=MARINE)
_sous_titre = Font(name="Calibri", size=10, color=GRIS)
_entete = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
_corps = Font(name="Calibri", size=10, color=ENCRE)
_accent = Font(name="Calibri", size=10, bold=True, color=ROSE)
_fond_entete = PatternFill("solid", fgColor=MARINE)
_fond_alterne = PatternFill("solid", fgColor=CARTE)
_bordure = Border(bottom=Side(style="thin", color=FILET))
_FORMAT_H = '[h]" h "mm'


def _h(td: timedelta | None):
    """Durée en fraction de jour (format Excel), pour permettre les calculs."""
    return td.total_seconds() / 86400 if td is not None else None


def _tableau(ws, ligne: int, entetes: list[str], lignes: list[list], formats: dict[int, str] | None = None,
             largeurs: list[int] | None = None) -> None:
    for c, titre in enumerate(entetes, 1):
        cell = ws.cell(ligne, c, titre)
        cell.font, cell.fill = _entete, _fond_entete
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[ligne].height = 30
    for i, valeurs in enumerate(lignes, 1):
        for c, v in enumerate(valeurs, 1):
            cell = ws.cell(ligne + i, c, v)
            cell.font, cell.border = _corps, _bordure
            if i % 2 == 0:
                cell.fill = _fond_alterne
            if formats and c in formats:
                cell.number_format = formats[c]
    ws.freeze_panes = ws.cell(ligne + 1, 1)
    if lignes:
        ws.auto_filter.ref = f"A{ligne}:{get_column_letter(len(entetes))}{ligne + len(lignes)}"
    for c, l in enumerate(largeurs or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = l


def _entete_feuille(ws, titre: str, debut: date, fin: date) -> None:
    ws["A1"] = titre
    ws["A1"].font = _titre
    ws["A2"] = f"Période du {format_date(debut)} au {format_date(fin)}"
    ws["A2"].font = _sous_titre
    ws.sheet_view.showGridLines = False


def exporter_excel(lignes: list[Synthese], debut: date, fin: date, mois: list[tuple[int, int]]) -> BytesIO:
    wb = Workbook()

    # --- Synthèse -----------------------------------------------------------
    ws = wb.active
    ws.title = "Synthèse"
    _entete_feuille(ws, "Badgeages : synthèse", debut, fin)
    entetes = ["Matricule", "Nom", "Prénom", "Établissement", "Manager", "Contrat hebdo (h)", "Jours travaillés",
               "Durée badgée", "Durée théorique", "Écart", "Alertes", "Saisies manuelles", "Hors zone"]
    multi_mois = len(mois) > 1
    if multi_mois:
        entetes += [libelle_mois(a, m).capitalize() for a, m in mois]
    donnees = []
    for s in lignes:
        u, c = s.utilisateur, s.compte_alertes()
        ligne = [u.matricule, u.nom, u.prenom, u.etablissement.nom if u.etablissement else "",
                 u.manager.nom_complet if u.manager else "", float(s.heures_contrat), s.jours_travailles,
                 _h(s.duree), _h(s.duree_attendue), _h(s.ecart), s.nb_alertes_graves, c.get("manuel", 0),
                 c.get("hors_zone", 0)]
        if multi_mois:
            pm = s.par_mois()
            ligne += [_h(pm.get(m, timedelta())) for m in mois]
        donnees.append(ligne)
    formats = {8: _FORMAT_H, 9: _FORMAT_H, 10: _FORMAT_H}
    formats.update({14 + i: _FORMAT_H for i in range(len(mois) if multi_mois else 0)})
    _tableau(ws, 4, entetes, donnees, formats, [12, 18, 16, 18, 22, 10, 10, 12, 12, 12, 9, 10, 9] + [12] * len(mois))
    note = ws.cell(len(donnees) + 6, 1,
                   "Durée théorique : contrat hebdomadaire ÷ 5 par jour ouvré (hors jours fériés) jusqu'à "
                   "aujourd'hui. Les absences (congés, maladie) ne sont pas prises en compte.")
    note.font = _sous_titre

    # --- Détail par journée -------------------------------------------------
    wd = wb.create_sheet("Détail")
    _entete_feuille(wd, "Badgeages : détail par journée", debut, fin)
    detail = []
    for s in lignes:
        u = s.utilisateur
        for j in s.journees_periode:
            for poste in j.postes:
                e, so = poste.entree, poste.sortie
                detail.append([
                    u.matricule, u.nom, u.prenom, j.jour, format_heure(poste.debut), format_heure(poste.fin),
                    _h(poste.duree),
                    _origine(e), ZONE_LIBELLES.get(e.zone_statut, "") if e else "",
                    _origine(so), ZONE_LIBELLES.get(so.zone_statut, "") if so else "",
                    "; ".join(x.motif for x in poste.pointages if x.motif) or "",
                    ", ".join(ALERTES[a] for a in j.alertes),
                ])
    _tableau(wd, 4, ["Matricule", "Nom", "Prénom", "Jour", "Arrivée", "Départ", "Durée", "Origine arrivée",
                     "Zone arrivée", "Origine départ", "Zone départ", "Motif", "Alertes de la journée"],
             detail, {4: "DD/MM/YYYY", 7: _FORMAT_H}, [12, 18, 16, 12, 9, 9, 10, 16, 16, 16, 16, 30, 40])

    # --- Légende ------------------------------------------------------------
    wl = wb.create_sheet("Légende")
    wl["A1"] = "Légende"
    wl["A1"].font = _titre
    lignes_legende = [
        ("Badge", "Badgeage réalisé en temps réel depuis l'application."),
        ("Saisie manuelle", "Déclaré a posteriori par le salarié (« J'ai oublié de badger »), avec motif."),
        ("Correction", "Ajouté ou modifié par un manager ou un administrateur (tracé dans le journal d'audit)."),
        ("Dans la zone", "Badgeage réalisé dans le périmètre d'un lieu autorisé (site ou domicile)."),
        ("Hors zone", "Badgeage réalisé en dehors de tout lieu autorisé."),
        ("Indéterminé", "Position trop imprécise pour conclure (fréquent sur ordinateur)."),
        ("Position refusée", "Le navigateur n'a pas transmis de position."),
        ("Confidentialité", "Aucune position GPS n'est enregistrée : seul le résultat du contrôle l'est."),
    ]
    for i, (k, v) in enumerate(lignes_legende, 3):
        wl.cell(i, 1, k).font = _accent
        wl.cell(i, 2, v).font = _corps
    wl.column_dimensions["A"].width = 20
    wl.column_dimensions["B"].width = 95
    wl.sheet_view.showGridLines = False

    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)
    return flux


def _origine(p) -> str:
    if p is None:
        return "Non badgé"
    libelle = {"badge": "Badge", "manuel": "Saisie manuelle", "correction": "Correction"}[p.source]
    if p.source == "badge" and p.appareil:
        libelle += " (smartphone)" if p.appareil == "mobile" else " (ordinateur)"
    if p.modifie_le:
        libelle += ", corrigé"
    if p.statut_validation == "en_attente":
        libelle += ", à valider"
    return libelle
